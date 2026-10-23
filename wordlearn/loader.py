import pandas as pd


def load_words(file_path: str) -> list[str]:
    """
    Load words from an Excel file.
    Assumes words are in the first column.
    """
    df = pd.read_excel(file_path)

    if df.empty:
        return []

    # Use the first column, drop blanks, and normalize spacing.
    words = (
        df.iloc[:, 0]
        .dropna()
        .astype(str)
        .str.strip()
        .tolist()
    )

    return words


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_WORDS = 100


def normalize_uploaded_words(values: list[str]) -> list[str]:
    words = []
    seen = set()
    for value in values:
        word = value.strip()
        if not word:
            continue
        if len(word) > 100:
            raise ValueError("单个单词或短语不能超过 100 个字符，请检查第一列。")
        if word.casefold() not in seen:
            seen.add(word.casefold())
            words.append(word)
    if not words:
        raise ValueError("词表为空，请在第一列填写单词，每行一个。")
    if len(words) > MAX_WORDS:
        raise ValueError(f"每次练习最多支持 {MAX_WORDS} 个不同单词，请缩小词表后重试。")
    return words


def parse_uploaded_words(content: bytes, filename: str) -> list[str]:
    import csv
    from io import BytesIO, StringIO
    from pathlib import Path
    from zipfile import ZipFile
    from openpyxl import load_workbook

    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".csv"}:
        raise ValueError("请上传 .xlsx 或 UTF-8 编码的 .csv 表格。")
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("文件不能为空且不能超过 5 MB。")
    values = []
    try:
        if suffix == ".csv":
            rows = csv.reader(StringIO(content.decode("utf-8-sig")))
            for row in rows:
                if row:
                    values.append(row[0])
        else:
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError("表格内容过大，请精简后上传。")
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                for row in workbook.worksheets[0].iter_rows(min_col=1, max_col=1, values_only=True):
                    if row[0] is not None:
                        values.append(str(row[0]))
            finally:
                workbook.close()
    except ValueError:
        raise
    except Exception:
        raise ValueError("无法读取表格，请确认文件未损坏，CSV 使用 UTF-8 编码。") from None
    values = [value.strip() for value in values if value.strip()]
    if values and values[0].casefold() in {"word", "words", "vocabulary", "单词", "英文", "英文单词"}:
        values.pop(0)
    return normalize_uploaded_words(values)
